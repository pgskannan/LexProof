"""Nebius Token Factory spike: can NVIDIA Nemotron run LexProof's real contract analysis?

Part A of claude/cursor-prompt-nebius-1-spike-and-provider.md. Script only; it
changes nothing in the app.

What it does
  1. GET /v1/models on Token Factory and lists every Nemotron model your key can use.
  2. For each Nemotron model (or --models), sends LexProof's *production* analysis
     prompt (services/analysis_prompt.py) for the chosen sample contracts.
     Mode a: response_format json_schema (strict) with ANALYSIS_RESPONSE_SCHEMA.
     Mode b (if a is rejected): response_format json_object + schema in the prompt.
  3. Validates every output against ANALYSIS_RESPONSE_SCHEMA with jsonschema.
  4. Optionally (--with-gemini) runs the same prompt through the app's own
     VertexGeminiProvider as the baseline.
  5. Prints a results table and writes spike_out/results.md, results.json and the
     raw output for each model/contract.

Run (PowerShell, from C:\\Projects\\LexProof):
  # one-time: put your key in backend\\.env (gitignored) as  NEBIUS_API_KEY=...
  .\\backend\\.venv\\Scripts\\python.exe ForNebius\\nebius_spike.py --with-gemini

Options:
  --contracts CONTRACT_01,CONTRACT_03   sample contract prefixes (default: both)
  --models id1,id2                      skip auto-discovery
  --list-only                           just print the Nemotron model IDs
  --with-gemini                         include the Vertex Gemini baseline

The API key is read from the environment / backend\\.env only and is never printed
or written to disk.
"""
from __future__ import annotations

import argparse
import ast
import asyncio
import importlib.util
import json
import os
import re
import sys
import time
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
BACKEND = ROOT / "backend"
SERVICES = BACKEND / "app" / "lexproof" / "services"
SAMPLES = ROOT / "sampleContracts"
OUT = Path(__file__).resolve().parent / "spike_out"
DEFAULT_BASE_URL = "https://api.tokenfactory.nebius.com/v1"

try:
    from dotenv import load_dotenv

    load_dotenv(BACKEND / ".env")
except ImportError:
    pass

import httpx  # noqa: E402  (in backend venv)
import jsonschema  # noqa: E402  (in backend venv)


# --------------------------------------------------------------------------- inputs
def load_analysis_prompt_module():
    spec = importlib.util.spec_from_file_location("lp_analysis_prompt", SERVICES / "analysis_prompt.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)  # dependency-free by design
    return module


def load_analysis_schema() -> dict:
    """Read ANALYSIS_RESPONSE_SCHEMA as a literal from vertex_ai.py (no app import)."""
    tree = ast.parse((SERVICES / "vertex_ai.py").read_text(encoding="utf-8"))
    for node in tree.body:
        if isinstance(node, ast.Assign) and any(getattr(t, "id", "") == "ANALYSIS_RESPONSE_SCHEMA" for t in node.targets):
            return ast.literal_eval(node.value)
    raise RuntimeError("ANALYSIS_RESPONSE_SCHEMA not found in vertex_ai.py")


def docx_text(path: Path) -> str:
    import docx

    document = docx.Document(str(path))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def find_contract(prefix: str) -> Path:
    matches = sorted(SAMPLES.glob(f"{prefix}*.docx"))
    if not matches:
        raise SystemExit(f"No sample contract matching {prefix}*.docx in {SAMPLES}")
    return matches[0]


# --------------------------------------------------------------------------- output parsing
THINK_RE = re.compile(r"<think>.*?</think>", re.DOTALL | re.IGNORECASE)


def clean_json_text(text: str) -> tuple[str, bool]:
    had_think = bool(THINK_RE.search(text or "")) or "</think>" in (text or "")
    text = THINK_RE.sub("", text or "")
    if "</think>" in text:  # unterminated opening tag stripped by the server
        text = text.split("</think>", 1)[1]
    text = text.strip()
    if text.startswith("```"):
        text = text.split("\n", 1)[-1]
        if text.rstrip().endswith("```"):
            text = text[: text.rstrip().rfind("```")]
    text = text.strip()
    if not text.startswith("{"):
        start, end = text.find("{"), text.rfind("}")
        if start != -1 and end > start:
            text = text[start : end + 1]
    return text, had_think


def validate(text: str, schema: dict) -> tuple[dict | None, str]:
    try:
        parsed = json.loads(text)
    except json.JSONDecodeError as exc:
        return None, f"invalid JSON: {exc.msg} at char {exc.pos}"
    errors = sorted(jsonschema.Draft202012Validator(schema).iter_errors(parsed), key=lambda e: list(e.path))
    if errors:
        first = errors[0]
        where = "/".join(str(p) for p in first.path) or "(root)"
        return parsed, f"{len(errors)} schema error(s); first at {where}: {first.message[:120]}"
    return parsed, ""


# --------------------------------------------------------------------------- Token Factory
class TokenFactory:
    def __init__(self, api_key: str, base_url: str):
        self.base_url = base_url.rstrip("/")
        self.client = httpx.Client(
            headers={"Authorization": f"Bearer {api_key}", "Content-Type": "application/json"},
            timeout=httpx.Timeout(300.0, connect=20.0),
        )

    def list_models(self) -> list[str]:
        response = self.client.get(f"{self.base_url}/models")
        response.raise_for_status()
        return sorted(m.get("id", "") for m in response.json().get("data", []))

    def chat(self, model: str, system: str, prompt: str, response_format: dict, max_tokens: int, no_think: bool = False) -> tuple[dict, int]:
        if no_think:
            system = system + " /no_think"
        body = {
            "model": model,
            "messages": [{"role": "system", "content": system}, {"role": "user", "content": prompt}],
            "temperature": 0.1,
            "max_tokens": max_tokens,
            "response_format": response_format,
        }
        if no_think:
            body["chat_template_kwargs"] = {"enable_thinking": False}
        for attempt, delay in enumerate((2, 6, 15, None)):
            started = time.monotonic()
            response = self.client.post(f"{self.base_url}/chat/completions", json=body)
            latency = int((time.monotonic() - started) * 1000)
            if response.status_code in (429, 500, 502, 503, 504) and delay is not None:
                time.sleep(delay)
                continue
            if response.status_code >= 400:
                raise httpx.HTTPStatusError(f"{response.status_code}: {response.text[:300]}", request=response.request, response=response)
            return response.json(), latency
        raise RuntimeError("unreachable")


def run_nemotron(tf: TokenFactory, model: str, system: str, prompt: str, schema: dict, max_tokens: int, no_think: bool = False) -> dict:
    modes = [
        ("a json_schema", {"type": "json_schema", "json_schema": {"name": "analysis", "schema": schema, "strict": True}}, prompt),
        (
            "b json_object",
            {"type": "json_object"},
            prompt + "\n\nReturn a single JSON object that conforms exactly to this JSON Schema:\n" + json.dumps(schema),
        ),
    ]
    notes = []
    for label, response_format, user_prompt in modes:
        try:
            payload, latency = tf.chat(model, system, user_prompt, response_format, max_tokens, no_think)
        except httpx.HTTPStatusError as exc:
            notes.append(f"{label} rejected: {str(exc)[:160]}")
            continue
        choice = (payload.get("choices") or [{}])[0]
        message = choice.get("message") or {}
        content = message.get("content") or ""
        reasoning_chars = len(message.get("reasoning_content") or message.get("reasoning") or "")
        text, had_think = clean_json_text(content)
        parsed, error = validate(text, schema)
        usage = payload.get("usage") or {}
        result = {
            "mode": label,
            "valid": not error,
            "error": error,
            "findings": len(parsed.get("findings", [])) if isinstance(parsed, dict) else 0,
            "risk_level": parsed.get("risk_level") if isinstance(parsed, dict) else None,
            "risk_score": parsed.get("risk_score") if isinstance(parsed, dict) else None,
            "latency_ms": latency,
            "prompt_tokens": usage.get("prompt_tokens"),
            "completion_tokens": usage.get("completion_tokens"),
            "finish_reason": choice.get("finish_reason"),
            "think_block": had_think or reasoning_chars > 0,
            "reasoning_chars": reasoning_chars,
            "raw": content,
            "parsed": parsed,
            "notes": notes,
        }
        if result["valid"] or label.startswith("b"):
            return result
        notes.append(f"{label} invalid (finish={choice.get('finish_reason')}, out_tokens={usage.get('completion_tokens')}, reasoning_chars={reasoning_chars}): {error}")
    return {"mode": "-", "valid": False, "error": "; ".join(notes), "findings": 0, "latency_ms": None, "raw": "", "parsed": None, "notes": notes}


# --------------------------------------------------------------------------- Gemini baseline
def run_gemini(system: str, prompt: str, schema: dict) -> dict:
    sys.path.insert(0, str(BACKEND / "app"))
    try:
        from lexproof.services.vertex_ai import VertexGeminiProvider  # type: ignore
    except Exception as exc:  # pragma: no cover - environment dependent
        return {"mode": "vertex", "valid": False, "error": f"import failed: {exc}", "findings": 0, "latency_ms": None, "raw": "", "parsed": None}

    class Request:
        model = None
        system_prompt = system

    Request.prompt = prompt
    provider = VertexGeminiProvider()
    try:
        response = asyncio.run(provider.complete(Request()))
    except Exception as exc:
        return {"mode": "vertex", "valid": False, "error": str(exc)[:200], "findings": 0, "latency_ms": None, "raw": "", "parsed": None}
    text, _ = clean_json_text(response.content)
    parsed, error = validate(text, schema)
    return {
        "mode": f"vertex {response.model}",
        "valid": not error,
        "error": error,
        "findings": len(parsed.get("findings", [])) if isinstance(parsed, dict) else 0,
        "risk_level": parsed.get("risk_level") if isinstance(parsed, dict) else None,
        "risk_score": parsed.get("risk_score") if isinstance(parsed, dict) else None,
        "latency_ms": response.latency_ms,
        "prompt_tokens": None,
        "completion_tokens": None,
        "finish_reason": None,
        "think_block": False,
        "raw": response.content,
        "parsed": parsed,
    }


# --------------------------------------------------------------------------- main
def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--contracts", default="CONTRACT_01,CONTRACT_03")
    parser.add_argument("--models", default="")
    parser.add_argument("--list-only", action="store_true")
    parser.add_argument("--tag", default="", help="suffix for results file names, e.g. run2")
    parser.add_argument("--with-gemini", action="store_true")
    parser.add_argument("--max-tokens", type=int, default=8192)
    parser.add_argument("--no-think", action="store_true", help="ask Nemotron to skip reasoning (chat_template_kwargs enable_thinking=false + /no_think)")
    args = parser.parse_args()

    api_key = os.getenv("NEBIUS_API_KEY", "").strip()
    if not api_key:
        print("NEBIUS_API_KEY is not set. Add it to backend\\.env (gitignored) or set it in this shell.")
        return 2
    tf = TokenFactory(api_key, os.getenv("NEBIUS_BASE_URL", DEFAULT_BASE_URL))

    all_models = tf.list_models()
    nemotron = [m for m in all_models if "nemotron" in m.lower()]
    print(f"Token Factory models visible to this key: {len(all_models)}; Nemotron: {len(nemotron)}")
    for m in nemotron:
        print(f"  - {m}")
    if args.list_only:
        return 0
    models = [m.strip() for m in args.models.split(",") if m.strip()] or [
        m for m in nemotron if not any(x in m.lower() for x in ("embed", "rerank", "reward", "guard", "safety", "parse", "vl", "omni"))
    ]
    if not models:
        print("No text Nemotron models found; pass --models explicitly.")
        return 3

    ap = load_analysis_prompt_module()
    schema = load_analysis_schema()
    system = ap.ANALYSIS_SYSTEM_PROMPT
    OUT.mkdir(parents=True, exist_ok=True)

    rows = []
    for prefix in [c.strip() for c in args.contracts.split(",") if c.strip()]:
        contract = find_contract(prefix)
        prompt = ap.build_analysis_prompt(docx_text(contract), "")  # playbook lives in Firestore; empty here
        print(f"\n== {contract.name} ({len(prompt):,} prompt chars)")
        targets = [("nemotron", m) for m in models] + ([("gemini", "vertex")] if args.with_gemini else [])
        for kind, model in targets:
            print(f"   {model} ...", flush=True)
            result = run_gemini(system, prompt, schema) if kind == "gemini" else run_nemotron(tf, model, system, prompt, schema, args.max_tokens, args.no_think)
            safe = re.sub(r"[^A-Za-z0-9._-]+", "_", model)
            (OUT / f"{prefix}__{safe}{('_' + args.tag) if args.tag else ''}.json").write_text(
                json.dumps({k: v for k, v in result.items() if k != "parsed"} | {"parsed": result.get("parsed")}, indent=2, ensure_ascii=False),
                encoding="utf-8",
            )
            rows.append({"contract": prefix, "model": result.get("mode", "").startswith("vertex") and result["mode"].split(" ", 1)[-1] or model, **{k: v for k, v in result.items() if k not in ("raw", "parsed")}})

    print(f"\nSettings: max_tokens={args.max_tokens} no_think={args.no_think}")
    header = "| Contract | Model | Mode | Schema-valid | Findings | Risk | Latency (ms) | Tokens in/out | Finish | <think> | Error |"
    lines = [header, "|" + "---|" * 11]
    for r in rows:
        tokens = f"{r.get('prompt_tokens') or '-'}/{r.get('completion_tokens') or '-'}"
        risk = f"{r.get('risk_level') or '-'} {r.get('risk_score') if r.get('risk_score') is not None else ''}".strip()
        lines.append(
            f"| {r['contract']} | {r['model']} | {r.get('mode')} | {'yes' if r['valid'] else 'NO'} | {r['findings']} | {risk} | "
            f"{r.get('latency_ms') or '-'} | {tokens} | {r.get('finish_reason') or '-'} | {'yes' if r.get('think_block') else 'no'} | {(r.get('error') or '')[:90]} |"
        )
    table = "\n".join(lines)
    print("\n" + table)
    tag = f"_{args.tag}" if args.tag else ""
    notes_md = "\n".join(f"- {r['contract']} {r['model']}: {n}" for r in rows for n in (r.get("notes") or []))
    (OUT / f"results{tag}.md").write_text(f"# Nebius spike results ({time.strftime('%Y-%m-%d %H:%M')}) max_tokens={args.max_tokens} no_think={args.no_think}\n\n{table}\n\n## Attempt notes\n{notes_md}\n", encoding="utf-8")
    (OUT / f"results{tag}.json").write_text(json.dumps(rows, indent=2), encoding="utf-8")
    print(f"\nSaved: {OUT / f'results{tag}.md'} and raw outputs in {OUT}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
